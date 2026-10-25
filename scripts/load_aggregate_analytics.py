#!/usr/bin/env python3

import os
import sys
import re
from pathlib import Path
from datetime import datetime, date

import psycopg
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DIR = ROOT / "data/raw/aggregate_analytics"


def integer(value):
    if value is None or str(value).strip() == "":
        return None
    return int(str(value).replace(",", "").strip())


def day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip(), "%m/%d/%Y").date()


def get_conn():
    import os
    import yaml
    from pathlib import Path

    profiles_path = Path.home() / ".dbt" / "profiles.yml"

    if not profiles_path.exists():
        raise FileNotFoundError(
            f"dbt profiles file not found: {profiles_path}"
        )

    with open(profiles_path, "r") as f:
        profiles = yaml.safe_load(f)

    profile = profiles["linkedin_intelligence"]

    target = profile.get("target", "dev")
    config = profile["outputs"][target]

    return psycopg.connect(
        host=config["host"],
        port=config.get("port", 5432),
        dbname=config["dbname"],
        user=config["user"],
        password=config["pass"],
        sslmode="require"
    )



DDL = """
CREATE SCHEMA IF NOT EXISTS raw;

CREATE TABLE IF NOT EXISTS raw.linkedin_aggregate_snapshots (
    snapshot_date date PRIMARY KEY,
    period_start date NOT NULL,
    period_end date NOT NULL,
    overall_impressions bigint,
    overall_reach bigint,
    total_followers bigint,
    source_file text NOT NULL,
    loaded_at timestamptz NOT NULL DEFAULT now()
);

CREATE TABLE IF NOT EXISTS raw.linkedin_daily_engagement (
    snapshot_date date NOT NULL,
    metric_date date NOT NULL,
    impressions bigint,
    engagements bigint,
    PRIMARY KEY (snapshot_date, metric_date)
);

CREATE TABLE IF NOT EXISTS raw.linkedin_follower_history (
    snapshot_date date NOT NULL,
    metric_date date NOT NULL,
    new_followers bigint,
    PRIMARY KEY (snapshot_date, metric_date)
);

CREATE TABLE IF NOT EXISTS raw.linkedin_audience_demographics (
    snapshot_date date NOT NULL,
    dimension text NOT NULL,
    category text NOT NULL,
    percentage_text text NOT NULL,
    percentage numeric(7,3),
    is_less_than boolean NOT NULL,
    PRIMARY KEY (snapshot_date, dimension, category)
);

CREATE TABLE IF NOT EXISTS raw.linkedin_content_demographics (
    snapshot_date date NOT NULL,
    dimension text NOT NULL,
    category text NOT NULL,
    percentage_text text NOT NULL,
    percentage numeric(7,3),
    is_less_than boolean NOT NULL,
    PRIMARY KEY (snapshot_date, dimension, category)
);
"""


def rows(sheet):
    return list(sheet.values)


def demographics(sheet):
    data = rows(sheet)

    for row in data[1:]:
        if not row[0] or not row[1] or not row[2]:
            continue

        raw = str(row[2]).strip()

        is_less_than = raw.startswith("<")

        percentage = (
            None
            if is_less_than
            else float(raw.replace("%", "").strip())
        )

        yield (
            str(row[0]).strip(),
            str(row[1]).strip(),
            raw,
            percentage,
            is_less_than,
        )


def process(path, conn):

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    required = {
        "DISCOVERY",
        "ENGAGEMENT",
        "FOLLOWERS",
        "AUDIENCE DEMOGRAPHICS",
        "CONTENT DEMOGRAPHICS",
    }

    missing = required - set(wb.sheetnames)

    if missing:
        raise ValueError(
            f"{path.name}: Missing worksheets: {missing}"
        )

    discovery_rows = rows(wb["DISCOVERY"])

    match = re.search(
        r"(\d{1,2}/\d{1,2}/\d{4})\s*-\s*"
        r"(\d{1,2}/\d{1,2}/\d{4})",
        str(discovery_rows[0][1])
    )

    if not match:
        raise ValueError("Cannot identify reporting dates")

    start = day(match.group(1))
    end = day(match.group(2))

    snapshot_date = end

    discovery = {
        str(row[0]).strip(): integer(row[1])
        for row in discovery_rows[1:]
        if row[0]
    }

    follower_data = rows(wb["FOLLOWERS"])

    total_followers = integer(
        follower_data[0][1]
    )

    engagement_rows = [
        (
            snapshot_date,
            day(row[0]),
            integer(row[1]),
            integer(row[2]),
        )
        for row in rows(wb["ENGAGEMENT"])[1:]
        if row[0]
    ]

    follower_rows = [
        (
            snapshot_date,
            day(row[0]),
            integer(row[1]),
        )
        for row in follower_data[3:]
        if row[0]
    ]

    # Validate reporting dates
    for row in engagement_rows + follower_rows:
        if not start <= row[1] <= end:
            raise ValueError(
                f"Unexpected date {row[1]} in {path.name}"
            )

    with conn.cursor() as cur:

        # Overall analytics snapshot
        cur.execute(
            """
            INSERT INTO raw.linkedin_aggregate_snapshots (
                snapshot_date,
                period_start,
                period_end,
                overall_impressions,
                overall_reach,
                total_followers,
                source_file
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s)

            ON CONFLICT (snapshot_date)
            DO UPDATE SET
                period_start = EXCLUDED.period_start,
                period_end = EXCLUDED.period_end,
                overall_impressions = EXCLUDED.overall_impressions,
                overall_reach = EXCLUDED.overall_reach,
                total_followers = EXCLUDED.total_followers,
                source_file = EXCLUDED.source_file,
                loaded_at = now()
            """,
            (
                snapshot_date,
                start,
                end,
                discovery.get("Impressions"),
                discovery.get("Members reached"),
                total_followers,
                path.name,
            )
        )

        # Refresh records belonging to this snapshot only
        tables = [
            "linkedin_daily_engagement",
            "linkedin_follower_history",
            "linkedin_audience_demographics",
            "linkedin_content_demographics",
        ]

        for table in tables:
            cur.execute(
                f"DELETE FROM raw.{table} "
                "WHERE snapshot_date = %s",
                (snapshot_date,)
            )

        # Daily engagement
        cur.executemany(
            """
            INSERT INTO raw.linkedin_daily_engagement
            (snapshot_date, metric_date, impressions, engagements)
            VALUES (%s, %s, %s, %s)
            """,
            engagement_rows
        )

        # Daily follower growth
        cur.executemany(
            """
            INSERT INTO raw.linkedin_follower_history
            (snapshot_date, metric_date, new_followers)
            VALUES (%s, %s, %s)
            """,
            follower_rows
        )

        # Audience demographics
        demographic_tables = [
            (
                "AUDIENCE DEMOGRAPHICS",
                "linkedin_audience_demographics",
            ),
            (
                "CONTENT DEMOGRAPHICS",
                "linkedin_content_demographics",
            ),
        ]

        for sheet_name, table in demographic_tables:

            values = [
                (snapshot_date, *record)
                for record in demographics(wb[sheet_name])
            ]

            cur.executemany(
                f"""
                INSERT INTO raw.{table} (
                    snapshot_date,
                    dimension,
                    category,
                    percentage_text,
                    percentage,
                    is_less_than
                )
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                values
            )

            print(f"{table}: {len(values)} records")

    print(f"\nProcessed: {path.name}")
    print(f"Snapshot date: {snapshot_date}")
    print(f"Engagement days: {len(engagement_rows)}")
    print(f"Follower days: {len(follower_rows)}")


def main():

    if len(sys.argv) > 1:
        paths = [
            Path(arg).expanduser().resolve()
            for arg in sys.argv[1:]
        ]
    else:
        paths = sorted(
            DEFAULT_DIR.glob("AggregateAnalytics_*.xlsx")
        )

    if not paths:
        raise SystemExit(
            f"No Excel files found in {DEFAULT_DIR}"
        )

    with get_conn() as conn:

        with conn.cursor() as cur:
            cur.execute(DDL)

        for path in paths:
            with conn.transaction():
                process(path, conn)

    print("\nLINKEDIN AGGREGATE ANALYTICS LOADED")
    print("-----------------------------------")
    print("Schema: raw")
    print("Existing post analytics: unchanged")
    print("Status: Complete")


if __name__ == "__main__":
    main()
